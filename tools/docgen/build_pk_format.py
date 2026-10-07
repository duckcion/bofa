"""Rebuild BOFA_Documentation.xlsx in the Platinum Kaizo docs layout.

Run after tools/docgen/generate_workbook.py:  python tools/docgen/build_pk_format.py
Writes development_reports/BOFA_Docs_PK_Format.xlsx. Needs the Platinum Kaizo docs (PK_DOCS)
for the list of later-gen moves to keep.
"""
import re
import sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import os
DECOMP = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
SRC = os.path.join(DECOMP, "development_reports", "BOFA_Documentation.xlsx")
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(DECOMP, "development_reports", "BOFA_Docs_PK_Format.xlsx")

def rd(sheet):
    return pd.read_excel(SRC, sheet_name=sheet, header=2)

def blank(v):
    return v is None or (isinstance(v, float) and pd.isna(v)) or str(v).strip() in ("", "nan")

def num(v):
    try:
        f = float(v)
        return int(f) if f == int(f) else f
    except (TypeError, ValueError):
        return v

# ---------------------------------------------------------------- styling
BG = "1F1F1F"
GOLD, YELLOW = "F1C232", "FFD966"
F_BG = PatternFill("solid", fgColor=BG)
F_NAME = PatternFill("solid", fgColor="141414")
F_MON = PatternFill("solid", fgColor="D9D9D9")
F_WHITE = PatternFill("solid", fgColor="FFFFFF")
F_ITEM = PatternFill("solid", fgColor="FFF2CC")
F_CHANGED = PatternFill("solid", fgColor="FF9900")
F_HEAD = PatternFill("solid", fgColor="434343")
F_ALT = PatternFill("solid", fgColor="F3F3F3")
THIN = Side(style="thin", color="7F7F7F")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=False)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=False)
WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)
FONT = "Arial"

TYPE_COLORS = {
    "Normal": "EFEFEF", "Fire": "EA9999", "Water": "9FC5E8", "Electric": "FFE599",
    "Grass": "93C47D", "Ice": "C9EEF5", "Fighting": "E06666", "Poison": "CB8BDD",
    "Ground": "F6B26B", "Flying": "96DADF", "Psychic": "F4A6C6", "Bug": "D8D179",
    "Rock": "D9C79E", "Ghost": "C3B6E3", "Dragon": "AFC0F5", "Dark": "BDB6B6",
    "Steel": "D0D0D0", "Fairy": "F7C6E0", "Stellar": "D5A6BD", "Mystery": "EFEFEF",
}
DARK_TYPES = set()  # all type colours are light enough for black text

def type_fill(t):
    return PatternFill("solid", fgColor=TYPE_COLORS.get(t, "EFEFEF"))

def type_font(t, **kw):
    return Font(name=FONT, color="FFFFFF" if t in DARK_TYPES else "000000", **kw)

def header_row(ws, row, headers, widths=None):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row, c, h)
        cell.font = Font(name=FONT, bold=True, color=YELLOW)
        cell.fill = PatternFill("solid", fgColor=BG)
        cell.alignment = CENTER
        cell.border = BOX
    if widths:
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = w

def body_cell(ws, r, c, v, fill=None, bold=False, align=CENTER, color="000000", italic=False):
    cell = ws.cell(r, c, None if blank(v) else num(v))
    cell.font = Font(name=FONT, bold=bold, color=color, italic=italic)
    cell.alignment = align
    cell.border = BOX
    if fill is not None:
        cell.fill = fill
    return cell

# ---------------------------------------------------------------- names
FORM_WORDS = {"Alola", "Galar", "Hisui", "Paldea", "Mega", "Gmax", "Primal"}

def form_label(species, const):
    base = "SPECIES_" + re.sub(r"[^A-Z0-9]", "_", species.upper())
    suffix = const[len(base):].strip("_") if const.startswith(base) else ""
    return f"{species} ({suffix.replace('_', ' ').title()})" if suffix else species

def showdown_id(species):
    words = species.replace("♀", "-F").replace("♂", "-M").split()
    if len(words) > 1 and words[-1] in FORM_WORDS:
        return re.sub(r"[^a-z0-9]", "", "".join(words[:-1]).lower()) + "-" + words[-1].lower()
    return re.sub(r"[^a-z0-9\-]", "", species.lower().replace(" ", ""))

def sprite_formula(species):
    url = f"https://play.pokemonshowdown.com/sprites/gen5/{showdown_id(species)}.png"
    return f'=IFERROR(_xlfn.IMAGE("{url}"),"")'

def pretty_loc(loc):
    if blank(loc) or str(loc).startswith("UNVERIFIED"):
        return "Unknown Location (no script reference)"
    parts = []
    for piece in str(loc).split(", "):
        segs = []
        for s in piece.split("_"):
            s = re.sub(r"(?<=[a-z])(?=[A-Z0-9])|(?<=[0-9])(?=[A-Z])", " ", s)
            s = re.sub(r"(\d) ([FR])\b", r"\1\2", s)
            s = re.sub(r"^SS ?", "S.S. ", s)
            segs.append(s)
        parts.append(" - ".join(segs))
    return " / ".join(parts)

# Rough Hoenn progression order for sorting trainer locations within a split.
MAP_ORDER = [
    "Route101", "OldaleTown", "Route103", "Route102", "PetalburgCity", "PetalburgCoast", "Route104",
    "PetalburgWoods", "RustboroCity", "Route116", "RusturfTunnel", "Hollowbrook", "GraniteShore",
    "Route105", "DewfordTown", "Route106", "GraniteCave", "Route107", "Route108", "Route109",
    "SlateportCity", "Route110", "MauvilleCity", "Route117", "VerdanturfTown", "TrainerGrove",
    "Route111", "Route112", "FieryPath", "Route113", "FallarborTown", "Route114", "MeteorFalls",
    "Route115", "MtChimney", "JaggedPass", "LavaridgeTown", "WraithwoodForest", "Route118",
    "Route119", "FortreeCity", "Route120", "Route121", "SafariZone", "LilycoveCity", "MtPyre",
    "Route122", "Route123", "MagmaHideout", "AquaHideout", "Route124", "MossdeepCity", "Route125",
    "ShoalCave", "Route126", "Route127", "Route128", "SeafloorCavern", "SootopolisCity",
    "CaveOfOrigin", "Route129", "Route130", "Route131", "PacifidlogTown", "Route132", "Route133",
    "Route134", "AbandonedShip", "SSTidal", "VictoryRoad", "EverGrandeCity", "BattleFrontier",
]
E4_ORDER = ["SidneysRoom", "PhoebesRoom", "GlaciasRoom", "DrakesRoom", "ChampionsRoom"]

def loc_key(loc):
    loc = "" if blank(loc) else str(loc)
    first = loc.split(", ")[0]
    base = first.split("_")[0]
    is_gym = "_Gym" in loc
    idx = next((i for i, m in enumerate(MAP_ORDER) if base.startswith(m)), len(MAP_ORDER))
    e4 = next((i for i, m in enumerate(E4_ORDER) if m in loc), -1)
    if loc.startswith("UNVERIFIED"):
        idx = 999
    nums = [int(n) for n in re.findall(r"\d+", first)]
    return (is_gym, idx, e4, nums, first)

# ---------------------------------------------------------------- load
# Cosmetic forms the source generator could not parse have no species name; skip them.
stats = rd("Pokemon Stats").dropna(subset=["Species"])
abil = rd("Ability Changes").dropna(subset=["Species"])
evo = rd("Evolution Changes").dropna(subset=["Species"])
learn = rd("Level-Up Learnsets").dropna(subset=["Species", "Level"])
moves = rd("Move Changes")

def mkey(name):
    return re.sub(r"[^a-z0-9]", "", str(name).lower())

move_type = {mkey(k): v for k, v in zip(moves["Move"], moves["New Type"])}
for _c in ["Orig Ability 1", "Orig Ability 2", "Orig Hidden", "New Ability 1", "New Ability 2", "New Hidden"]:
    abil[_c] = abil[_c].apply(lambda v: "-" if blank(v) or str(v).strip() == "None" else v)
abil_by_const = abil.set_index("Const")

# ---------------------------------------------------------------- Gen 1-6 scope
LAST_GEN6_DEX = 721       # Volcanion
LAST_GEN6_MOVE = 621      # MOVE_HYPERSPACE_FURY (MOVES_COUNT_GEN6 = 622)
# Forms of Gen 1-6 species that were only introduced in Gen 7+.
LATER_FORMS = re.compile(
    r"ALOLA|GALAR|HISUI|PALDEA|GMAX|TOTEM|PIKACHU_(HOENN|KALOS|ORIGINAL|PARTNER|SINNOH|STARTER|UNOVA|WORLD)"
    r"|EEVEE_STARTER|DIALGA_ORIGIN|PALKIA_ORIGIN|BASCULIN_WHITE|GRENINJA_(ASH|BATTLE_BOND)|POWER_CONSTRUCT"
    r"|ZYGARDE_COMPLETE")

def gen6_species(dex, const):
    return not blank(dex) and float(dex) <= LAST_GEN6_DEX and not LATER_FORMS.search(str(const))

GEN6_CONSTS = {c for d, c in zip(stats["Nat Dex #"], stats["Form/Const"]) if gen6_species(d, c)}
GEN6_NAMES = {form_label(sp, c).upper() for sp, c in zip(stats["Species"], stats["Form/Const"]) if c in GEN6_CONSTS}
stats = stats[stats["Form/Const"].isin(GEN6_CONSTS)]
abil = abil[abil["Const"].isin(GEN6_CONSTS)]
evo = evo[evo["Const"].isin(GEN6_CONSTS)]

move_ids = {}
with open(DECOMP + r"\include\constants\moves.h", encoding="utf-8") as fh:
    for m in re.finditer(r"#define (MOVE_[A-Z0-9_]+)\s+(\d+)\b", fh.read()):
        move_ids[m.group(1)] = int(m.group(2))
# Later-gen moves are kept only if Platinum Kaizo's docs also use them.
PK_DOCS = r"C:\Users\garet\Downloads\Copy of Platinum Kaizo Docs.xlsx"
PK_MOVES = {mkey(n) for n in pd.read_excel(PK_DOCS, sheet_name="Moves")["Name"].dropna()}
moves = moves[(moves["Const"].map(lambda c: move_ids.get(c, 9999)) <= LAST_GEN6_MOVE)
              | moves["Move"].map(mkey).isin(PK_MOVES)]
GEN6_MOVES = {mkey(m) for m in moves["Move"]}

# learnsets keyed by symbol (in file order), plus species-name -> symbol for trainer default moves
learnsets = {}
const_symbol = {}
sym_owner = {}
for _, r in learn.iterrows():
    sym = r["Learnset Symbol"]
    const_symbol.setdefault(r["Const"], sym)
    # Forms sharing a learnset repeat its rows once per form; keep only the first form's copy.
    if sym_owner.setdefault(sym, r["Const"]) == r["Const"]:
        learnsets.setdefault(sym, []).append((int(r["Level"]), r["Move"]))

def species_const(name):
    return "SPECIES_" + re.sub(r"[^A-Z0-9]+", "_", name.upper()).strip("_")

def default_moves(species, level):
    """Expansion fills empty trainer movesets with the last 4 distinct level-up moves <= level."""
    sym = const_symbol.get(species_const(species))
    if sym is None:
        return []
    known = []
    for lv, mv in learnsets[sym]:
        if lv <= level and mv not in known:
            if len(known) == 4:
                known.pop(0)
            known.append(mv)
    return known

def default_ability(species):
    c = species_const(species)
    if c in abil_by_const.index:
        return abil_by_const.loc[c, "New Ability 1"]
    return ""

wb = Workbook()
wb.remove(wb.active)

# ================================================================ Personal
ws = wb.create_sheet("Personal")
heads = [" ", "Name", "HP", "Attack", "Defense", "Sp. Atk", "Sp. Def", "Speed", "BST",
         "Type 1", "Type 2", "Ability 1", "Ability 2", "Hidden Ability"]
header_row(ws, 1, heads, [6, 24, 9, 9, 9, 9, 9, 9, 10, 10, 10, 16, 16, 16])
STAT_COLS = [("HP", 3), ("Atk", 4), ("Def", 5), ("SpA", 6), ("SpD", 7), ("Spe", 8), ("BST", 9)]
r = 2
for _, s in stats.iterrows():
    body_cell(ws, r, 1, s["Nat Dex #"])
    body_cell(ws, r, 2, form_label(s["Species"], s["Form/Const"]).upper(), bold=True, align=LEFT)
    for key, col in STAT_COLS:
        old, new = num(s[f"Orig {key}"]), num(s[f"New {key}"])
        if not blank(old) and not blank(new) and old != new:
            diff = new - old
            body_cell(ws, r, col, f"{new} ({'+' if diff > 0 else ''}{diff})", fill=F_CHANGED, bold=True)
        else:
            body_cell(ws, r, col, new)
    old_t = [t.strip() for t in str(s["Orig Typing"]).split("/")]
    new_t = [t.strip() for t in str(s["New Typing"]).split("/")]
    new_t += new_t[:1] * (2 - len(new_t))
    old_t += old_t[:1] * (2 - len(old_t))
    for i, t in enumerate(new_t[:2]):
        c = body_cell(ws, r, 10 + i, t, fill=type_fill(t))
        c.font = type_font(t, bold=t != old_t[i])
        if t != old_t[i]:
            c.value = f"{t} (was {old_t[i]})"
            c.border = Border(left=Side("medium", "FF9900"), right=Side("medium", "FF9900"),
                              top=Side("medium", "FF9900"), bottom=Side("medium", "FF9900"))
    if s["Form/Const"] in abil_by_const.index:
        a = abil_by_const.loc[s["Form/Const"]]
        for i, k in enumerate(["1", "2"]):
            ch = a[f"Orig Ability {k}"] != a[f"New Ability {k}"]
            body_cell(ws, r, 12 + i, a[f"New Ability {k}"], fill=F_CHANGED if ch else None, bold=ch)
        ch = a["Orig Hidden"] != a["New Hidden"]
        body_cell(ws, r, 14, a["New Hidden"], fill=F_CHANGED if ch else None, bold=ch)
    r += 1
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:N{r-1}"

# ================================================================ Ability Changes
ws = wb.create_sheet("Ability Changes")
header_row(ws, 1, ["Pokémon Name", "Ability 1", "Ability 2", "Hidden Ability", "Previously"],
           [26, 18, 18, 18, 46])
r = 2
for _, a in abil[abil["Changed?"].notna()].iterrows():
    body_cell(ws, r, 1, form_label(a["Species"], a["Const"]).upper(), bold=True, align=LEFT)
    for i, k in enumerate(["Ability 1", "Ability 2", "Hidden"]):
        ch = a[f"Orig {k}"] != a[f"New {k}"]
        body_cell(ws, r, 2 + i, a[f"New {k}"], fill=F_CHANGED if ch else None, bold=ch)
    body_cell(ws, r, 5, f"{a['Orig Ability 1']} / {a['Orig Ability 2']} / {a['Orig Hidden']}",
              align=LEFT, color="666666", italic=True)
    r += 1
ws.freeze_panes = "B2"

# ================================================================ Evolutions
def split_evos(text):
    out = []
    if blank(text) or str(text).strip() == "None":
        return out
    for part in str(text).split("; "):
        if " -> " in part:
            how, res = part.rsplit(" -> ", 1)
        else:
            how, res = part, ""
        if how == "None":
            continue
        m = re.match(r"Lv (\d+)$", how)
        if m:
            out.append(("Level Up", m.group(1), res))
            continue
        m = re.match(r"(.+?) \((.+)\)$", how)
        if m:
            out.append((m.group(1), m.group(2), res))
            continue
        m = re.match(r"(Specific Map|Friendship Move Type|Move|Level Female|Level Male|Party|Item Hold Day|Item Hold Night) (.+)", how)
        if m:
            out.append((m.group(1), m.group(2).replace("MAP_", "").replace("TYPE_", "").replace("_", " ").title(), res))
            continue
        out.append((how, "", res))
    return out

ws = wb.create_sheet("Evolutions")
max_e = max(len(split_evos(v)) for v in evo["New Evolution"])
heads = ["Name"]
for i in range(max_e):
    sfx = "" if i == 0 else f" {i+1}"
    heads += [f"Method{sfx}", f"Required{sfx}", f"Result{sfx}"]
heads += ["Previously"]
header_row(ws, 1, heads, [24] + [14, 16, 16] * max_e + [60])
r = 2
for _, e in evo.iterrows():
    ev = split_evos(e["New Evolution"])
    if not ev:
        continue
    changed = not blank(e["Changed?"])
    old = set(split_evos(e["Orig Evolution"]))
    body_cell(ws, r, 1, form_label(e["Species"], e["Const"]).upper(), bold=True, align=LEFT,
              fill=F_CHANGED if changed else None)
    for i, (m, req, res) in enumerate(ev):
        hl = F_CHANGED if changed and (m, req, res) not in old else None
        body_cell(ws, r, 2 + i * 3, m, fill=hl)
        body_cell(ws, r, 3 + i * 3, req, fill=hl)
        body_cell(ws, r, 4 + i * 3, res.upper(), fill=hl)
    if changed:
        body_cell(ws, r, 2 + max_e * 3, e["Orig Evolution"], align=LEFT, color="666666", italic=True)
    r += 1
ws.freeze_panes = "B2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(heads))}{r-1}"

# ================================================================ Level-Up Learnsets
ws = wb.create_sheet("Level-Up Learnsets")
seen_sym = []
sym_name = {}
for _, l in learn.iterrows():
    sym = l["Learnset Symbol"]
    if sym not in sym_name and sym_owner[sym] in GEN6_CONSTS:
        sym_name[sym] = form_label(l["Species"], l["Const"])
        seen_sym.append(sym)
# Only Gen 1-6 moves go in the grid; any later-gen moves still in the game's learnset are listed at the end.
g6_sets = {s: [(lv, mv) for lv, mv in learnsets[s] if mkey(mv) in GEN6_MOVES] for s in seen_sym}
later = {s: [mv for lv, mv in learnsets[s] if mkey(mv) not in GEN6_MOVES] for s in seen_sym}
max_m = max(len(v) for v in g6_sets.values())
heads = ["Name"]
for i in range(max_m):
    sfx = "" if i == 0 else f" {i+1}"
    heads += [f"Move{sfx}", f"Lv{sfx}"]
heads += ["Other Gen 7+ moves still in learnset"]
header_row(ws, 1, heads)
ws.column_dimensions["A"].width = 26
for i in range(max_m):
    ws.column_dimensions[get_column_letter(2 + i * 2)].width = 17
    ws.column_dimensions[get_column_letter(3 + i * 2)].width = 5
ws.column_dimensions[get_column_letter(len(heads))].width = 60
PLAIN = Alignment(horizontal="center", vertical="center", wrap_text=False)
PLAIN_L = Alignment(horizontal="left", vertical="center", wrap_text=False)
F_LV = PatternFill("solid", fgColor="F3F3F3")
r = 2
for sym in seen_sym:
    body_cell(ws, r, 1, sym_name[sym].upper(), bold=True, align=PLAIN_L)
    for i, (lv, mv) in enumerate(g6_sets[sym]):
        body_cell(ws, r, 2 + i * 2, mv, align=PLAIN)
        body_cell(ws, r, 3 + i * 2, "Evo" if lv == 0 else lv, align=PLAIN, fill=F_LV)
    body_cell(ws, r, len(heads), ", ".join(later[sym]), align=PLAIN_L, color="666666")
    r += 1
ws.freeze_panes = "B2"

# ================================================================ Moves (current data)
ws = wb.create_sheet("Moves")
heads = ["Name", "Type", "Category", "Power", "Accuracy", "PP", "Effect", "Effect Chance (%)", "Secondary Effect"]
header_row(ws, 1, heads, [26, 11, 11, 8, 10, 6, 22, 10, 30])
r = 2
for _, m in moves.iterrows():
    changed = not blank(m["Changed?"])
    body_cell(ws, r, 1, m["Move"], bold=True, align=LEFT, fill=F_CHANGED if changed else None)
    c = body_cell(ws, r, 2, m["New Type"], fill=type_fill(m["New Type"]))
    c.font = type_font(m["New Type"])
    body_cell(ws, r, 3, m["New Category"])
    body_cell(ws, r, 4, "-" if num(m["New Power"]) in (0, 1) and m["New Category"] == "Status" else m["New Power"])
    body_cell(ws, r, 5, "-" if num(m["New Acc"]) == 0 else m["New Acc"])
    body_cell(ws, r, 6, m["New PP"])
    body_cell(ws, r, 7, m["New Effect"])
    body_cell(ws, r, 8, m["Secondary Chance"])
    sec = "" if blank(m["Secondary Effects"]) else str(m["Secondary Effects"]).replace("MOVE_EFFECT_", "").replace("_", " ").title()
    body_cell(ws, r, 9, sec, align=LEFT)
    r += 1
ws.freeze_panes = "B2"
ws.auto_filter.ref = f"A1:I{r-1}"

# ================================================================ Move Changes (text list, like PK)
ws = wb.create_sheet("Move Changes")
ws.column_dimensions["A"].width = 130
c = ws.cell(1, 1, "Move Changes (Gen 1-6 moves + Platinum Kaizo extras)")
c.font = Font(name=FONT, bold=True, size=14, color=YELLOW)
c.fill = F_BG
FIELDS = [("Power", " bp"), ("Acc", " acc"), ("PP", " PP"), ("Type", " type"), ("Category", ""), ("Priority", " priority"),
          ("Target", "")]
r = 2
for _, m in moves[moves["Changed?"].notna()].iterrows():
    diffs = []
    for f, unit in FIELDS:
        o, n = num(m[f"Orig {f}"]), num(m[f"New {f}"])
        if not blank(o) and o != n:
            diffs.append(f"{o}{unit} -> {n}{unit}")
    new_sec = ", ".join(str(x) for x in (m["Secondary Effects"], m["Secondary Chance"]) if not blank(x))
    mech = (m["Orig Effect"] != m["New Effect"] or str(m["Orig Secondary"] if not blank(m["Orig Secondary"]) else "") != new_sec
            or not blank(m["Recoil %"]))
    if not diffs and not mech:
        continue
    head = f"{m['Move']} ({m['Description']})" if mech and not blank(m["Description"]) else f"{m['Move']}"
    c = ws.cell(r, 1, head + (": " + ", ".join(diffs) if diffs else ""))
    c.font = Font(name=FONT)
    if r % 2:
        c.fill = F_ALT
    c.border = BOX
    c.alignment = LEFT
    r += 1

# ================================================================ Trainer split sheets
split_sheets = ["Split 1", "Split 2", "Split 3", "Split 4", "Split 5", "Split 6", "Split 7",
                "Split 8", "Elite Four & Champion", "Postgame", "Unassigned", "Boss Battles"]

def load_trainers(sheet):
    df = rd(sheet)
    trainers, cur = [], None
    for _, row in df.iterrows():
        if not blank(row["Trainer"]):
            cur = {k: row[k] for k in ["Trainer", "Class", "Trainer ID", "Location", "Split",
                                       "Mandatory?", "Format", "AI Flags", "Trainer Items"]}
            cur["mons"] = []
            trainers.append(cur)
        cur["mons"].append(row)
    return trainers

ALL_TRAINER_NAMES = {}   # trainer ID -> name, to tell real rematches (same name) from numbered IDs
for _s in ["Split 1", "Split 2", "Split 3", "Split 4", "Split 5", "Split 6", "Split 7", "Split 8",
           "Elite Four & Champion", "Postgame", "Unassigned"]:
    _d = rd(_s).dropna(subset=["Trainer ID"])
    ALL_TRAINER_NAMES.update(zip(_d["Trainer ID"].astype(str), _d["Trainer"].astype(str)))

_REMATCH_TIER = {}
for _m in re.finditer(r"REMATCH\(([^)]*)\)", open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..",
                                                          "src", "battle_setup.c"), encoding="utf-8").read()):
    _ids = [x.strip() for x in _m.group(1).split(",")][:-1]
    for _k, _tid in enumerate(_ids[1:], 1):
        if _tid != _ids[0]:
            _REMATCH_TIER.setdefault(_tid, _k)

def trainer_title(t):
    name = str(t["Trainer"]).title()
    cls = "" if blank(t["Class"]) else str(t["Class"])
    title = f"{cls} {name}".strip()
    tid = str(t["Trainer ID"])
    if tid in _REMATCH_TIER:          # only trainers in the game's rematch table
        title += f" (Rematch {_REMATCH_TIER[tid]})"
    return title

def trainer_subtitle(t):
    bits = []
    fmt = str(t["Format"])
    if fmt == "Double":
        bits.append("DOUBLE BATTLE")
    elif fmt.startswith("Tag"):
        bits.append("TAG BATTLE" + (" (with partner)" if "partner" in fmt else " (two trainers)"))
    if t.get("Weather"):
        bits.append(str(t["Weather"]))
    mand = str(t["Mandatory?"])
    if mand.startswith("REQUIRED"):
        bits.append("Mandatory")
    elif mand.startswith("Forces"):
        bits.append("Forced on sight" + (" " + re.search(r"\(.*\)", mand).group(0) if "(" in mand else ""))
    elif mand.startswith("Opt-in"):
        bits.append("Optional")
    if not blank(t["Trainer Items"]):
        bits.append("Items: " + str(t["Trainer Items"]))
    return "\n".join(bits)

F_LOC = PatternFill("solid", fgColor="D9D9D9")
F_BAR = PatternFill("solid", fgColor="434343")
INK = "000000"
MUTED = "666666"

def box_pair(ws, r, col):
    """Border a 2-wide merged cell on its outer edges only."""
    ws.cell(r, col).border = Border(left=THIN, top=THIN, bottom=THIN)
    ws.cell(r, col + 1).border = Border(right=THIN, top=THIN, bottom=THIN)

def label(ws, r, text):
    c = ws.cell(r, 17, text)
    c.font = Font(name=FONT, bold=True, color=INK)
    c.alignment = CENTER

def write_split(name, trainers):
    ws = wb.create_sheet(name[:31])
    ws.sheet_view.showGridLines = True
    widths = {"A": 2.25, "B": 7, "C": 7, "D": 7, "Q": 16, "R": 2.75}
    for col in range(1, 19):
        L = get_column_letter(col)
        ws.column_dimensions[L].width = widths.get(L, 12.5)

    ws.merge_cells("A1:D3")
    ws.merge_cells("E1:Q1")
    ws.row_dimensions[1].height = 50
    c = ws.cell(1, 5, f"BOFA - {name.upper()}")
    c.font = Font(name=FONT, size=28, bold=True, color=INK)
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells("E2:Q2")
    c = ws.cell(2, 5, f"{len(trainers)} trainers.  Move cells are shaded by move type.  "
                      "Trainers with no custom moveset use their last 4 level-up moves.")
    c.font = Font(name=FONT, size=10, color=MUTED)
    c.alignment = LEFT
    ws.merge_cells("E3:Q3")
    c = ws.cell(3, 5, "Every trainer Pokémon has a role-based nature; ability is slot 1 unless the trainer file sets it.")
    c.font = Font(name=FONT, size=10, color=MUTED)
    c.alignment = LEFT

    r = 5
    last_loc = None
    ordered = sorted(trainers, key=lambda t: (loc_key(t["Location"]),
                                              "Leader" in str(t["Class"]) or "Champion" in str(t["Class"]),
                                              str(t["Trainer ID"])))
    if trainers and all("Order" in t for t in trainers):
        ordered = sorted(trainers, key=lambda t: t["Order"])
    if name == "Boss Battles":
        ordered = sorted(trainers, key=lambda t: (int(re.sub(r"\D", "", str(t["Split"])) or 99)
                                                  if str(t["Split"]).startswith("Split") else 98,
                                                  loc_key(t["Location"]), str(t["Trainer ID"])))
    for t in ordered:
        loc = pretty_loc(t["Location"])
        if loc != last_loc:
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=17)
            c = ws.cell(r, 2, loc + (f"   ({t['Split']})" if name == "Boss Battles" else ""))
            c.font = Font(name=FONT, size=14, bold=True, color=INK)
            c.alignment = CENTER
            for cc in range(2, 18):
                ws.cell(r, cc).fill = F_LOC
            ws.row_dimensions[r].height = 24
            r += 2
            last_loc = loc
        mons = t["mons"]
        # name bar
        ws.row_dimensions[r].height = 20
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        c = ws.cell(r, 2, str(t["Class"]))
        c.font = Font(name=FONT, size=9, color="FFFFFF")
        c.alignment = CENTER
        ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=17)
        c = ws.cell(r, 5, trainer_title(t))
        c.font = Font(name=FONT, size=12, bold=True, color="FFFFFF")
        c.alignment = CENTER
        for cc in range(2, 18):
            ws.cell(r, cc).fill = F_BAR
        sub = trainer_subtitle(t)
        # sprite row
        ws.row_dimensions[r + 1].height = 69
        for i, m in enumerate(mons[:6]):
            col = 5 + i * 2
            ws.merge_cells(start_row=r + 1, start_column=col, end_row=r + 1, end_column=col + 1)
            ws.cell(r + 1, col, sprite_formula(str(m["Species"])))
            ws.cell(r + 1, col).alignment = CENTER
        if sub:
            c = ws.cell(r + 1, 17, sub)
            c.font = Font(name=FONT, size=9, color=INK)
            c.alignment = WRAP
        # info rows
        R_MON, R_NAT, R_ITEM, R_MV = r + 2, r + 3, r + 4, r + 5
        label(ws, R_MON, "Pokémon")
        label(ws, R_NAT, "Nature / Ability")
        label(ws, R_ITEM, "Item")
        label(ws, R_MV, "Moves")
        ws.merge_cells(start_row=R_MV, start_column=17, end_row=R_MV + 3, end_column=17)
        ws.merge_cells(start_row=R_ITEM, start_column=2, end_row=R_ITEM, end_column=4)
        c = ws.cell(R_ITEM, 2, "AI Flags:")
        c.font = Font(name=FONT, bold=True, color=INK)
        c.alignment = CENTER
        ws.merge_cells(start_row=R_MV, start_column=2, end_row=R_MV + 3, end_column=4)
        c = ws.cell(R_MV, 2, "" if blank(t["AI Flags"]) else str(t["AI Flags"]).replace(", ", "\n"))
        c.font = Font(name=FONT, size=9, color=INK)
        c.alignment = WRAP

        for i, m in enumerate(mons[:6]):
            col = 5 + i * 2
            sp = str(m["Species"])
            lv = num(m["Level"])
            for rr in (R_MON, R_ITEM, R_MV, R_MV + 1, R_MV + 2, R_MV + 3):
                ws.merge_cells(start_row=rr, start_column=col, end_row=rr, end_column=col + 1)
            c = ws.cell(R_MON, col, f"Lv {lv} {sp}")
            c.font = Font(name=FONT, bold=True)
            c.fill = F_MON
            c.alignment = CENTER
            box_pair(ws, R_MON, col)
            nat = str(m["Nature"])
            nat = "Hardy" if blank(m["Nature"]) or "default" in nat else nat
            ab = str(m["Ability"])
            ab = default_ability(sp) if blank(m["Ability"]) or "default" in ab else ab
            for k, v in enumerate([nat, ab]):
                c = ws.cell(R_NAT, col + k, v)
                c.font = Font(name=FONT, size=9)
                c.alignment = CENTER
                c.border = BOX
            item = "(None)" if blank(m["Held Item"]) else str(m["Held Item"])
            c = ws.cell(R_ITEM, col, item)
            c.font = Font(name=FONT)
            c.fill = F_ITEM
            c.alignment = CENTER
            box_pair(ws, R_ITEM, col)
            mv_list = [m[f"Move {k}"] for k in range(1, 5) if not blank(m[f"Move {k}"])]
            if not mv_list:
                mv_list = default_moves(sp, int(lv) if isinstance(lv, (int, float)) else 100)
            for k in range(4):
                mv = mv_list[k] if k < len(mv_list) else ""
                c = ws.cell(R_MV + k, col, mv or "-")
                c.alignment = CENTER
                box_pair(ws, R_MV + k, col)
                c.font = Font(name=FONT, color=INK)
                if mv:
                    c.fill = type_fill(move_type.get(mkey(mv), "Normal"))
        r = R_MV + 5
        if t.get("Note"):
            ws.merge_cells(start_row=r - 1, start_column=5, end_row=r - 1, end_column=17)
            c = ws.cell(r - 1, 5, "Note: " + str(t["Note"]))
            c.font = Font(name=FONT, size=9, italic=True, color="C00000" if "NEEDS DECISION" in str(t["Note"]) else MUTED)
            c.alignment = LEFT
            r += 1
    ws.freeze_panes = "A4"
    return ws

# ---------------------------------------------------------------- Split 1 restructure plan
import split1_plan as plan

loaded = {s: load_trainers(s) for s in split_sheets}
by_id = {}
for s in split_sheets:
    if s != "Boss Battles":
        for t in loaded[s]:
            by_id.setdefault(str(t["Trainer ID"]), t)

def plan_mon(m):
    row = {"Species": m["species"], "Level": m["level"], "Ability": m["ability"] or "(default slot 1)",
           "Nature": m["nature"] or "(Hardy default)", "Held Item": m["item"]}
    for k in range(4):
        row[f"Move {k+1}"] = m["moves"][k] if k < len(m["moves"]) else None
    return row

def plan_trainer(p):
    if "from_id" in p:
        src = by_id[p["from_id"]]
        t = dict(src)
        mons = [dict(m) for m in src["mons"]]
        for m, lv in zip(mons, p.get("levels", [])):
            m["Level"] = lv
        t["mons"] = mons
    else:
        t = {"Trainer": p["name"], "Class": p["cls"], "Trainer ID": p["id"], "Format": "Single",
             "AI Flags": "Smart Trainer", "Trainer Items": None,
             "mons": [plan_mon(m) for m in p["mons"]]}
    t.update({"Location": p["loc"], "Split": "Split 1", "Order": p["order"], "Note": p.get("note")})
    t["Mandatory?"] = {"REQUIRED": "REQUIRED", "Opt-in": "Opt-in"}.get(p["status"], t.get("Mandatory?", p["status"]))
    return t

split1 = [plan_trainer(p) for p in plan.SPLIT1_TRAINERS]
moved_ids = {p["from_id"] for p in plan.SPLIT1_TRAINERS if "from_id" in p}
moved_ids |= set(plan.REMOVED_FROM_MAP)   # no longer on any map
route103 = []
for t in loaded["Split 1"]:
    tid = str(t["Trainer ID"])
    if tid in plan.ROUTE_103_IDS:
        t = dict(t)
        t["Split"] = "Unassigned"
        t["Note"] = "Moved out of Split 1 (Route 103, across the water). Assign to a later split."
        route103.append(t)
loaded["Split 1"] = split1
for s in split_sheets:
    if s not in ("Split 1", "Boss Battles"):
        loaded[s] = [t for t in loaded[s] if str(t["Trainer ID"]) not in moved_ids]
loaded["Unassigned"] = route103 + loaded["Unassigned"]

for t in loaded["Split 2"]:
    tid = str(t["Trainer ID"])
    if tid in plan.SPLIT2_ORDER:
        t["Note"] = plan.SPLIT2_NOTES.get(tid)
# Split 2 in play order: the listed trainers first, then everyone else by location
loaded["Split 2"] = sorted(loaded["Split 2"], key=lambda t: (plan.SPLIT2_ORDER.get(str(t["Trainer ID"]), 99),
                                                            loc_key(t["Location"]), str(t["Trainer ID"])))
for i, t in enumerate(loaded["Split 2"]):
    t["Order"] = i

# ---- Splits 1-3: exactly the trainers on the maps, in fight order (split1_plan.FIGHT_ORDER)
import json as _json
def map_battle_weather(folder):
    try:
        w = _json.load(open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "data", "maps",
                                         folder, "map.json"), encoding="utf-8")).get("weather", "")
    except OSError:
        return None
    if "FOG" in w: return None   # fog is visual only (B_OVERWORLD_FOG = GEN_7)
    if "RAIN" in w or "THUNDER" in w: return "Map weather: Rain"
    if "DROUGHT" in w: return "Map weather: Harsh sunlight"
    if "SANDSTORM" in w: return "Map weather: Sandstorm"
    if "SNOW" in w: return "Map weather: Snow"
    return None
placed = set()
for sname, order in plan.FIGHT_ORDER.items():
    out = []
    for i, (tid, folder, status, fmt, note) in enumerate(order):
        src = by_id.get(tid)
        if src is None:
            print("FIGHT_ORDER: no data for", tid); continue
        t = dict(src)
        t.update({"Location": folder, "Split": sname, "Order": i, "Note": note, "Format": fmt,
                  "Mandatory?": status, "Weather": map_battle_weather(folder), "Trainer Items": None})
        out.append(t); placed.add(tid)
    loaded[sname] = out
for sname in split_sheets:
    if sname not in plan.FIGHT_ORDER and sname != "Boss Battles":
        loaded[sname] = [t for t in loaded[sname]
                         if str(t["Trainer ID"]) not in placed and str(t["Trainer ID"]) not in plan.REMOVED_FROM_MAP]

for s in split_sheets:
    write_split(s, loaded[s])

# ================================================================ Encounters (grid per map/method)
wild = rd("Wild Encounters")
wild["Map"] = wild["Map"].fillna("MAP_NO_MAP_ASSIGNED")
for mp, sp in plan.ENCOUNTER_SPLIT.items():
    wild.loc[wild["Map"] == mp, "Split"] = sp
# per-map land rates (BOFA engine override), read from the same header the game uses
CUSTOM_LAND = {m.group(1): [int(v) for v in m.group(2).split(",")]
               for m in re.finditer(r"\{\s*(MAP_\w+),\s*\{([\d,\s]+)\}\s*\}",
                                    open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..",
                                                      "src", "data", "wild_encounters_custom_rates.h"), encoding="utf-8").read())}
RATES = {
    "Land": [20, 20, 10, 10, 10, 10, 5, 5, 4, 4, 1, 1],
    "Water": [60, 30, 5, 4, 1],
    "Rock Smash": [60, 30, 5, 4, 1],
    "Fishing": [70, 30, 60, 20, 20, 40, 40, 15, 4, 1],
}
ROD = lambda slot: "Old Rod" if slot < 2 else ("Good Rod" if slot < 5 else "Super Rod")
ws = wb.create_sheet("Encounters")
ws.column_dimensions["A"].width = 30
ws.column_dimensions["B"].width = 10
ws.column_dimensions["C"].width = 12
ws.column_dimensions["D"].width = 6
r = 1
groups = {}
order = []
seen_slots = set()
for _, w in wild.iterrows():
    method = w["Method"]
    slot = int(w["Slot"])
    # unused tables (no map) and Altering Cave's extra Mystery Event tables never appear in game:
    # only the first table for a map is live
    if w["Map"] == "MAP_NO_MAP_ASSIGNED" or (w["Map"], method, slot) in seen_slots:
        continue
    seen_slots.add((w["Map"], method, slot))
    sub = ROD(slot) if method == "Fishing" else method
    key = (w["Map"], sub) if w["Map"] != "MAP_NO_MAP_ASSIGNED" else (w["Map"], sub, w["Encounter Rate"])
    if key not in groups:
        groups[key] = {"split": w["Split"], "rate": w["Encounter Rate"], "mons": {}}
        order.append(key)
    rates = CUSTOM_LAND.get(w["Map"], RATES["Land"]) if method == "Land" else RATES.get(method, [])
    pct = rates[slot] if slot < len(rates) else 0
    g = groups[key]["mons"].setdefault(w["Species"], [pct and 0, 999, 0])
    g[0] += pct
    g[1] = min(g[1], int(w["Min Level"]))
    g[2] = max(g[2], int(w["Max Level"]))
# Vertical layout (like the PK docs): each location is a 3-column block (% | Pokemon | Lv), one
# encounter per row; methods run down the left and line up across the locations in a band.
METHOD_ORDER = ["Land", "Water", "Good Rod", "Super Rod", "Rock Smash"]  # no Old Rod in BOFA
METHOD_LABEL = {"Land": "Grass / Cave", "Water": "Surf"}
METHOD_FILL = {m: PatternFill("solid", fgColor=c) for m, c in {
    "Land": "D9EAD3", "Water": "CFE2F3", "Good Rod": "DDE7F0", "Super Rod": "B7CDE8", "Rock Smash": "EAD9C2"}.items()}
METHOD_HEAD = {m: PatternFill("solid", fgColor=c) for m, c in {
    "Land": "93C47D", "Water": "9FC5E8", "Good Rod": "A9BED3", "Super Rod": "6FA8DC", "Rock Smash": "C9B07A"}.items()}
PER_BAND = 10000  # all locations side by side in one row; only methods stack
maps = []
for key in order:
    if key[0] not in [m for m, _ in maps]:
        maps.append((key[0], groups[key]["split"]))
split_rank = lambda sp: int(re.search(r"\d+", str(sp)).group()) if re.search(r"\d+", str(sp)) else 99
maps.sort(key=lambda m: (split_rank(m[1]), [k[0] for k in order].index(m[0])))
by_map = {}
for key in order:
    by_map.setdefault(key[0], {})[key[1]] = groups[key]

ws.delete_rows(1, ws.max_row)
ws.column_dimensions["A"].width = 14
for band_start in range(0, len(maps), PER_BAND):
    band = maps[band_start:band_start + PER_BAND]
    for k, (mp, split) in enumerate(band):
        c0 = 2 + k * 4
        ws.column_dimensions[get_column_letter(c0)].width = 7
        ws.column_dimensions[get_column_letter(c0 + 1)].width = 16
        ws.column_dimensions[get_column_letter(c0 + 2)].width = 8
        ws.column_dimensions[get_column_letter(c0 + 3)].width = 2
        loc = pretty_loc(mp.replace("MAP_", "").title().replace("_", "")) if mp != "MAP_NO_MAP_ASSIGNED" else "No map assigned"
        ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + 2)
        body_cell(ws, r, c0, loc, fill=F_MON, bold=True)
        ws.merge_cells(start_row=r + 1, start_column=c0, end_row=r + 1, end_column=c0 + 2)
        body_cell(ws, r + 1, c0, split)
        for j, h in enumerate(["%", "Pokemon", "Lv"]):
            body_cell(ws, r + 2, c0 + j, h, bold=True, fill=F_ALT)
    body_cell(ws, r, 1, "Location", bold=True, fill=F_MON)
    body_cell(ws, r + 1, 1, "Split")
    body_cell(ws, r + 2, 1, "", fill=F_ALT)
    r += 3
    for method in METHOD_ORDER:
        tables = [by_map.get(mp, {}).get(method) for mp, _ in band]
        if not any(tables):
            continue
        height = max(len(t["mons"]) for t in tables if t)
        ws.merge_cells(start_row=r, start_column=1, end_row=r + height, end_column=1)
        body_cell(ws, r, 1, METHOD_LABEL.get(method, method), bold=True, align=WRAP, fill=METHOD_HEAD[method])
        for k, t in enumerate(tables):
            c0 = 2 + k * 4
            ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + 2)
            body_cell(ws, r, c0, f"Encounter rate {t['rate']}" if t else "", italic=True, fill=METHOD_HEAD[method])
            mons = sorted(t["mons"].items(), key=lambda kv: -kv[1][0]) if t else []
            for i in range(height):
                if i < len(mons):
                    sp, (pc, lo, hi) = mons[i]
                    body_cell(ws, r + 1 + i, c0, f"{pc}%", fill=METHOD_FILL[method])
                    body_cell(ws, r + 1 + i, c0 + 1, sp, align=LEFT, fill=METHOD_FILL[method])
                    body_cell(ws, r + 1 + i, c0 + 2, f"{lo}" if lo == hi else f"{lo}-{hi}", fill=METHOD_FILL[method])
                else:
                    for j in range(3):
                        body_cell(ws, r + 1 + i, c0 + j, "", fill=METHOD_FILL[method])
        r += height + 1
    r += 2
ws.freeze_panes = "B1"

# ================================================================ plain restyled tables
def clean_method(m):
    return "Mart" if str(m).startswith("Purchased") else str(m).replace(" (visible)", "")

def plain_sheet(src, title, drop=(), widths=None, keep=None, df=None):
    df = rd(src) if df is None else df
    if keep is not None:
        df = df[keep(df)]
    if "Method" in df.columns and src in ("Item Locations", "TM & HM Locations"):
        df = df.assign(Method=df["Method"].map(clean_method), Map=df["Map"].map(pretty_loc))
    df = df[[c for c in df.columns if c not in drop and not c.startswith("Unnamed")]]
    ws = wb.create_sheet(title)
    header_row(ws, 1, list(df.columns), widths)
    if not widths:
        for i, c in enumerate(df.columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = max(10, min(45, len(c) + 4,
                max([len(str(v)) for v in df[c].head(200)] + [0]) + 2))
    for rr, row in enumerate(df.itertuples(index=False), 2):
        fill = F_ALT if rr % 2 else None
        for cc, v in enumerate(row, 1):
            body_cell(ws, rr, cc, v, fill=fill, align=LEFT)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}{len(df)+1}"

def split_fixups(df):
    if "Coords" in df.columns:
        ys = df["Coords"].astype(str).str.extract(r"\(\s*\d+,\s*(\d+)\)")[0].astype(float)
        south = (df["Map"] == "Route104") & (ys >= 41)
        df.loc[south, "Split"] = "Split 1"
        xs = df["Coords"].astype(str).str.extract(r"\(\s*(\d+),")[0].astype(float)
        df = df.assign(Note=df["Note"] if "Note" in df.columns else "")
        cut = (df["Map"] == "PetalburgWoods") & (((xs == 35) & (ys == 20)) | ((xs == 42) & (ys == 20)) | ((xs == 45) & (ys == 7)))
        df.loc[cut, "Note"] = "Behind Cut trees - not reachable in Split 2"
        later = ((df["Map"] == "Route116") & (xs >= 60)) | ((df["Map"] == "RusturfTunnel") & (ys >= 10))
        df.loc[later, "Note"] = "Past Rusturf Tunnel - not reachable in Split 2"
        surf = (df["Map"] == "Route103") & (xs >= 30)
        df.loc[surf, "Note"] = "Across the water - needs Surf"
    if "Item" in df.columns and "Method" in df.columns:
        late = df["Method"].astype(str).str.startswith("Purchased") & df["Item"].isin(["Great Ball", "Timer Ball"])
        df = df.assign(Note=df.get("Note", ""))
        df.loc[late, "Note"] = "Only sold after Badge 2"
    return df

SPLIT_RANK = {**{f"Split {i}": i for i in range(1, 9)}, "Elite Four & Champion": 9, "Postgame": 10, "Unassigned": 11}
METHOD_RANK = {"Item Ball": 0, "Hidden": 1, "Mart": 3, "Purchased": 3}

def chrono_sort(df):
    """Order rows the way a player meets them: split, then visit order of the map, then pickups before marts."""
    def map_rank(split, mp):
        order = plan.CHRONO_MAPS.get(split, [])
        mp = str(mp)
        if split in ("Split 2",) and mp.startswith("RustboroCity_Gym"):
            return 90        # the gym comes last in Split 2
        hits = [(len(prefix), i) for i, prefix in enumerate(order) if mp.startswith(prefix)]
        if hits:
            return max(hits)[1]          # longest matching name wins (PetalburgCity_Gym over PetalburgCity)
        return 100 + loc_key(mp)[1]
    keys = [(SPLIT_RANK.get(str(r["Split"]), 12), map_rank(str(r["Split"]), r["Map"]),
             METHOD_RANK.get(str(r["Method"]).split(" (")[0], 2), str(r.get("Item", r.get("TM/HM", ""))))
            for _, r in df.iterrows()]
    df = df.assign(_k=keys).sort_values("_k", kind="stable").drop(columns="_k")
    return df.reset_index(drop=True)

def planned_item_locs():
    df = split_fixups(rd("Item Locations"))
    for item, mp in plan.REMOVED_ITEMS:
        df = df[~((df["Item"] == item) & (df["Map"] == mp))]
    add = pd.DataFrame([{"Item": it, "Method": me, "Map": mp, "Split": sp, "Note": note}
                        for it, me, mp, sp, note in plan.NEW_ITEMS])
    return chrono_sort(pd.concat([add, df], ignore_index=True))

def planned_tm_locs():
    df = split_fixups(rd("TM & HM Locations").assign(Note=""))
    add = pd.DataFrame([{"TM/HM": t, "Move Taught": mv, "Method": how, "Map": mp, "Split": sp, "Note": note}
                        for t, mv, how, mp, sp, note in (
        ("HM06", "Rock Smash", "NPC gift (end of forest)", "ViridianForest", "Split 1", "Shares its flag with the Mauville gift"),
        ("TM17", "Protect", "Gym reward (Norman), x1", "PetalburgCity_Gym", "Split 1", "Only single-copy TM"),
        ("TM09", "Bullet Seed", "NPC gift (boy), x2", "Route104", "Split 1", "North end of Route 104 South"),
        ("HM01", "Cut", "NPC gift (the Cutter, in his house)", "RustboroCity_CuttersHouse", "Split 2", "Usable with Badge 1"),
        ("TM68", "Wing Attack", "NPC gift (Trainer School student), x2", "RustboroCity_PokemonSchool", "Split 2", "New BOFA TM"),
        ("TM39", "Rock Tomb", "Gym reward (Roxanne), x2", "RustboroCity_Gym", "Split 2", ""))])
    return chrono_sort(pd.concat([add, df], ignore_index=True))

def planned_gifts():
    df = rd("Gifts & Trades")
    _, species = plan.GIFTS["remove_trade_species_at"]
    # the old Viridian Forest starter trades: their trade.h entries are unused now
    df = df[~((df["Type"] == "In-game Trade") & (df["Species"].isin(species)))]
    df = df.drop_duplicates(subset=["Type", "Species", "Map"]).copy()
    johto = df["Type"].str.startswith("Gift") & df["Species"].isin(species)
    room = johto & (df["Map"] == "StarterRoom")
    df.loc[room, "Moves / Notes"] = plan.GIFTS["johto_gift_note"]
    df.loc[room, "Split"] = "Split 1"
    df.loc[johto & ~room, "Moves / Notes"] = "Postgame: Birch's National Dex Johto-starter choice"
    df.loc[johto & ~room, "Split"] = "Postgame"
    return df

plain_sheet("Progression", "Progression", drop=("Status",))

prog = pd.DataFrame(plan.PROGRESSION, columns=["Step", "Area", "What Happens", "Trainers", "Gate / Event", "Status"])
plain_sheet(None, "Split 1 Progression", df=prog, widths=[6, 34, 60, 46, 60, 34])
prog2 = pd.DataFrame(plan.SPLIT2_PROGRESSION, columns=["Step", "Area", "What Happens", "Trainers", "Gate / Event", "Status"])
plain_sheet(None, "Split 2 Progression", df=prog2, widths=[6, 30, 60, 60, 70, 28])
prog3 = pd.DataFrame(plan.SPLIT3_PROGRESSION, columns=["Step", "Area", "What Happens", "Trainers", "Gate / Event", "Status"])
plain_sheet(None, "Split 3 Progression", df=prog3, widths=[6, 30, 60, 52, 56, 40])

# Split 1 plan check: counts audit + contradictions
ws = wb.create_sheet("Split 1 Plan Check")
header_row(ws, 1, ["Check", "Result"], [36, 120])
req = [t for t in split1 if t["Mandatory?"] == "REQUIRED"]
rival_variants = [t for t in split1 if t["Class"] == "Rival"]
by_loc = {}
for t in split1:
    by_loc.setdefault(pretty_loc(t["Location"]), []).append(t)
rows = [
    ("Mandatory battles a player fights", f"{len(req) - len(rival_variants) + 1} (approved: 16) "
     f"(rival counted once; {len(rival_variants)} rival variants for Brendan/May x 3 starters)"),
    ("Optional trainers", ", ".join(trainer_title(t) for t in split1 if t["Mandatory?"] == "Opt-in")),
] + [(f"Trainers at {loc}", f"{len(ts)}: " + ", ".join(trainer_title(t) for t in ts)) for loc, ts in by_loc.items()] + [
    ("Route 103 trainers moved out", f"{len(route103)} moved to the Unassigned sheet; Route 103 rival battle removed"),
    ("Gen 2 starter", "Gift in the Littleroot lab kept; Viridian Forest starter trades removed (Gifts & Trades)"),
    ("Items", "Everstone added at Petalburg Coast (replaces Oval Stone); HM06 Rock Smash added at Viridian Forest"),
] + [("FLAG: " + k, v) for k, v in plan.FLAGS]
for rr, (k, v) in enumerate(rows, 2):
    body_cell(ws, rr, 1, k, bold=True, align=LEFT, color="C00000" if k.startswith("FLAG") else "000000")
    c = body_cell(ws, rr, 2, v, align=Alignment(horizontal="left", vertical="center", wrap_text=True))
ws.freeze_panes = "A2"
plain_sheet("Pokemon Availability", "Pokemon Availability", drop=("Status",),
            keep=lambda df: df["Species"].astype(str).str.upper().isin(GEN6_NAMES))
# Gift Pokemon, laid out like Item Locations: grouped by split and location, one row per Pokemon.
ws = wb.create_sheet("Gift Pokemon")
GIFT_HEAD = ["Pokemon", "How you get it", "You give", "Lv", "Nature", "Ability", "IVs", "Held Item", "Moves", "Notes", "Obtained?"]
header_row(ws, 1, GIFT_HEAD, [26, 38, 22, 10, 10, 34, 12, 12, 36, 40, 11])
rr = 2
last_split = last_loc = None
for split, loc, how, give, mon, lv, nature, abil, ivs, item, moves, note in plan.GIFT_POKEMON:
    if split != last_split:
        for c in range(1, len(GIFT_HEAD) + 1): body_cell(ws, rr, c, split if c == 1 else "", fill=PatternFill("solid", fgColor="D9D9D9"), align=LEFT, bold=(c == 1))
        rr += 1; last_split, last_loc = split, None
    if loc != last_loc:
        for c in range(1, len(GIFT_HEAD) + 1): body_cell(ws, rr, c, loc if c == 1 else "", fill=F_ALT, align=LEFT)
        rr += 1; last_loc = loc
    for c, val in enumerate([mon, how, give, lv, nature, abil, ivs, item, moves, note, "☐"], 1):
        body_cell(ws, rr, c, val, align=LEFT if c in (1, 2, 3, 6, 9, 10) else CENTER)
    rr += 1
ws.freeze_panes = "A2"
plain_sheet("TM & HM Locations", "TM & HM Locations", drop=("Status", "Item Const", "Flag"), df=planned_tm_locs())

# ================================================================ Items (Gen 1-6 only)
# Items introduced in Gen 7 or later (by constant name).
LATER_ITEMS = re.compile(
    r"^(STRANGE_BALL|BEAST_BALL|MAX_HONEY|PEWTER_CRUNCHIES|BIG_MALASADA|ABILITY_PATCH|.*_MINT|EXP_CANDY_.*"
    r"|DYNAMAX_CANDY|LURE|SUPER_LURE|MAX_LURE|MAX_MUSHROOMS|BOTTLE_CAP|GOLD_BOTTLE_CAP|FOSSILIZED_.*"
    r"|GALARICA_.*|ARMORITE_ORE|DYNITE_ORE|ICE_STONE|SWEET_APPLE|TART_APPLE|CRACKED_POT|CHIPPED_POT"
    r"|STRAWBERRY_SWEET|LOVE_SWEET|BERRY_SWEET|CLOVER_SWEET|FLOWER_SWEET|STAR_SWEET|RIBBON_SWEET"
    r"|.*_NECTAR|.*_MEMORY|RUSTED_SWORD|RUSTED_SHIELD|.*IUM_Z|ELECTRIC_SEED|PSYCHIC_SEED|MISTY_SEED"
    r"|GRASSY_SEED|ADRENALINE_ORB|TERRAIN_EXTENDER|PROTECTIVE_PADS|THROAT_SPRAY|EJECT_PACK"
    r"|HEAVY_DUTY_BOOTS|BLUNDER_POLICY|ROOM_SERVICE|UTILITY_UMBRELLA|CATCHING_CHARM|EXP_CHARM"
    r"|ROTOM_CATALOG|ZYGARDE_CUBE|N_SOLARIZER|N_LUNARIZER|REINS_OF_UNITY|Z_POWER_RING|DYNAMAX_BAND"
    r"|ABILITY_SHIELD|CLEAR_AMULET|PUNCHING_GLOVE|COVERT_CLOAK|LOADED_DICE|AUSPICIOUS_ARMOR"
    r"|BOOSTER_ENERGY|BIG_BAMBOO_SHOOT|GIMMIGHOUL_COIN|LEADERS_CREST|MALICIOUS_ARMOR|MIRROR_HERB"
    r"|SCROLL_OF_.*|TERA_ORB|TINY_BAMBOO_SHOOT|.*_TERA_SHARD|ADAMANT_CRYSTAL|GRISEOUS_CORE"
    r"|LUSTROUS_GLOBE|BLACK_AUGURITE|LINKING_CORD|PEAT_BLOCK|BERSERK_GENE|FAIRY_FEATHER|SYRUPY_APPLE"
    r"|UNREMARKABLE_TEACUP|MASTERPIECE_TEACUP|.*_MASK|.*_MOCHI|GLIMMERING_CHARM|METAL_ALLOY"
    r"|JUBILIFE_MUFFIN|REMEDY|FINE_REMEDY|SUPERB_REMEDY|AUX_.*|CHOICE_DUMPLING|SWAP_SNACK"
    r"|TWICE_SPICED_RADISH|POKESHI_DOLL|TM\d+|NONE)$")
POCKETS = {"POCKET_ITEMS": "Items", "POCKET_POKE_BALLS": "Poké Balls", "POCKET_TM_HM": "TMs & HMs",
           "POCKET_BERRIES": "Berries", "POCKET_KEY_ITEMS": "Key Items"}

GEN_LATEST = 9

def eval_config(expr):
    """Resolve config ternaries like '(I_PRICE >= GEN_7) ? 200 : 300' with every config at GEN_LATEST."""
    expr = expr.strip()
    if "?" not in expr:
        return num(expr.strip("() "))
    cond, rest = expr.split("?", 1)
    a, b = rest.split(":", 1)
    cond = re.sub(r"\bI_\w+\b", str(GEN_LATEST), cond)
    cond = re.sub(r"\bGEN_LATEST\b", str(GEN_LATEST), cond)
    cond = re.sub(r"\bGEN_(\d+)\b", r"\1", cond).replace("&&", " and ").replace("||", " or ")
    try:
        ok = eval(cond, {"__builtins__": {}})
    except Exception:
        ok = True
    return eval_config(a) if ok else eval_config(b)

first_branch = eval_config


with open(DECOMP + r"\src\data\items.h", encoding="utf-8") as fh:
    src = fh.read()
items = []
for m in re.finditer(r"^    \[ITEM_([A-Z0-9_]+)\] =\s*\{(.*?)^    \},", src, re.M | re.S):
    const, body = m.group(1), m.group(2)
    if LATER_ITEMS.match(const):
        continue
    name = re.search(r'\.name = (?:_|ITEM_NAME)\("(.*?)"\)', body)
    if not name:
        continue
    price = re.search(r"\.price = (.*?),\n", body)
    pocket = re.search(r"\.pocket = (\w+)", body)
    fling = re.search(r"\.flingPower = (\d+)", body)
    desc = re.search(r"\.description = COMPOUND_STRING\((.*?)\),\n", body, re.S)
    desc = " ".join(re.findall(r'"(.*?)"', desc.group(1))).replace("\\n", " ").replace("  ", " ") if desc else ""
    items.append({"name": name.group(1), "pocket": POCKETS.get(pocket.group(1) if pocket else "", ""),
                  "price": first_branch(price.group(1)) if price else 0,
                  "fling": int(fling.group(1)) if fling else "", "desc": desc})

locs = planned_item_locs()
found = {}
for _, l in locs.iterrows():
    if blank(l["Item"]):
        continue
    key = "HM06" if str(l["Item"]) == "HM Rock Smash" else str(l["Item"])
    found.setdefault(key, []).append(f"{pretty_loc(l['Map'])} ({l['Split']}, {clean_method(l['Method'])})")

ws = wb.create_sheet("Items")
header_row(ws, 1, ["Name", "Pocket", "Price", "Fling Power", "Description", "Where to Find"],
           [22, 12, 9, 11, 70, 80])
order = {"Items": 0, "Poké Balls": 1, "Berries": 2, "TMs & HMs": 3, "Key Items": 4, "": 5}
items.sort(key=lambda it: (order.get(it["pocket"], 5),))
for rr, it in enumerate(items, 2):
    fill = F_ALT if rr % 2 else None
    body_cell(ws, rr, 1, it["name"], bold=True, align=LEFT, fill=fill)
    body_cell(ws, rr, 2, it["pocket"], fill=fill)
    body_cell(ws, rr, 3, it["price"] or "-", fill=fill)
    body_cell(ws, rr, 4, it["fling"] or "-", fill=fill)
    body_cell(ws, rr, 5, it["desc"], align=LEFT, fill=fill)
    body_cell(ws, rr, 6, "; ".join(found.get(it["name"], [])), align=LEFT, fill=fill)
ws.freeze_panes = "B2"
ws.auto_filter.ref = f"A1:F{len(items)+1}"

# ---- Item Locations in the Platinum Kaizo checklist format:
#      area header rows in play order; Item | Location | Requirements | Obtained?
import glob, json
REPO_ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..")
def ball_quantities():
    q = {}
    for f in glob.glob(os.path.join(REPO_ROOT, "data", "maps", "*", "map.json")):
        mp = os.path.basename(os.path.dirname(f))
        try:
            j = json.load(open(f, encoding="utf-8"))
        except Exception:
            continue
        for o in j.get("object_events", []):
            if o.get("script") == "Common_EventScript_FindItem":
                q[(mp, o["x"], o["y"])] = max(1, int(o.get("movement_range_x") or 1))
    return q

def requirement_from_note(note):
    note = "" if pd.isna(note) else str(note)
    low = note.lower()
    if "surf" in low: return "Surf"
    if "cut tree" in low: return "Cut"
    if "rusturf tunnel" in low: return "After Rusturf Tunnel"
    if "badge" in low: return note.replace("Only sold after ", "")
    return ""

def location_text(method, coords, note):
    method = str(method); coords = "" if pd.isna(coords) or str(coords) in ("-", "nan") else str(coords)
    note = "" if pd.isna(note) else str(note)
    m = re.match(r"(NPC gift|NPC choice|Gym reward|Dropped by)\s*\(?(.*?)\)?$", method)
    if method.startswith("Item Ball"): text = f"Item ball at {coords}" if coords else "Item ball"
    elif method.startswith("Hidden"): text = f"At {coords} (hidden)" if coords else "(hidden)"
    elif method.startswith("NPC gift"): text = "From " + method[len("NPC gift"):].strip(" ()")
    elif method.startswith("Gym reward"): text = "From " + method[len("Gym reward"):].strip(" ()")
    elif method.startswith("NPC choice"): text = "NPC choice " + method[len("NPC choice"):].strip()
    elif method.startswith("Berry"): text = f"Berry tree at {coords}" if coords else "Berry tree"
    else: text = method + (f" at {coords}" if coords else "")
    if note and not requirement_from_note(note) and note not in text and "flag" not in note.lower():
        text += f" ({note})"
    return text

_TM_ORDER = re.findall(r"F\((\w+)\)", open(os.path.join(REPO_ROOT, "include", "constants", "tms_hms.h"), encoding="utf-8").read())
_HM_START = _TM_ORDER.index("CUT") if "CUT" in _TM_ORDER else len(_TM_ORDER)
def tm_name(name):
    """'TM Bullet Seed x2' / 'TM68 Wing Attack x2' / 'HM Rock Smash' -> 'TM09 - Bullet Seed x2' style."""
    m = re.match(r"^(TM|HM)(\d*)\s+(?!-)(.+?)(\s+x\d+)?$", name)
    if not m: return name
    kind, num, move, qty = m.group(1), m.group(2), m.group(3), m.group(4) or ""
    key = re.sub(r"[^A-Z0-9]+", "_", move.upper()).strip("_")
    if not num and key in _TM_ORDER:
        i = _TM_ORDER.index(key)
        num = f"{i + 1:02d}" if i < _HM_START else f"{i - _HM_START + 1:02d}"
        kind = "TM" if i < _HM_START else "HM"
    return f"{kind}{num} - {move}{qty}"

def area_of(mp):
    return pretty_loc(str(mp)).split(" - ")[0]

items = planned_item_locs()
items = items[items["Item"].astype(str) != "0"]
tms = planned_tm_locs()
tms = tms[~tms["Method"].astype(str).str.startswith(("NPC gift", "Gym reward"))]
tms = tms.assign(Item=tms["TM/HM"].astype(str) + " - " + tms["Move Taught"].astype(str))[["Item", "Method", "Map", "Split", "Note", "Coords"]]
allrows = chrono_sort(pd.concat([items, tms], ignore_index=True))
qty = ball_quantities()

rows = []      # (kind, values)
cur_split = cur_area = None
marts = {}
for _, r in allrows.iterrows():
    split, mp = str(r["Split"]), r["Map"]
    area = area_of(mp)
    if split != cur_split:
        rows.append(("split", split)); cur_split = split; cur_area = None
    if area != cur_area:
        rows.append(("area", area)); cur_area = area
    method = str(r["Method"])
    if method.startswith(("Mart", "Purchased")):
        key = (split, area, pretty_loc(str(mp)))
        if key not in marts:
            marts[key] = []
            rows.append(("mart", key))
        marts[key].append(str(r["Item"]) + (" (after Badge 2)" if "Badge 2" in str(r.get("Note", "")) else ""))
        continue
    name = tm_name(str(r["Item"]))
    m = re.match(r"\((\d+),\s*(\d+)\)", str(r.get("Coords", "")))
    if m and method.startswith("Item Ball"):
        n = qty.get((str(mp), int(m.group(1)), int(m.group(2))), 1)
        if n > 1 and not re.search(r"x\d+$", name): name += f" x{n}"
    rows.append(("item", [name, location_text(method, r.get("Coords"), r.get("Note")), requirement_from_note(r.get("Note")), "☐"]))

ws = wb.create_sheet("Item Locations")
header_row(ws, 1, ["Item", "Location", "Requirements", "Obtained?"], [30, 62, 24, 11])
F_SPLIT = PatternFill("solid", fgColor="D9D9D9")
F_AREA = PatternFill("solid", fgColor="F3F3F3")
rr = 2
for kind, v in rows:
    if kind == "split":
        for c in range(1, 5): body_cell(ws, rr, c, v if c == 1 else "", fill=F_SPLIT, align=LEFT)
    elif kind == "area":
        for c in range(1, 5): body_cell(ws, rr, c, v if c == 1 else "", fill=F_AREA, align=LEFT)
    elif kind == "mart":
        split, area, where = v
        body_cell(ws, rr, 1, "Mart stock", align=LEFT)
        body_cell(ws, rr, 2, f"{where}: " + ", ".join(marts[v]), align=LEFT)
        body_cell(ws, rr, 3, "", align=LEFT); body_cell(ws, rr, 4, "", align=LEFT)
    else:
        for c, val in enumerate(v, 1): body_cell(ws, rr, c, val, align=LEFT)
    rr += 1
ws.freeze_panes = "A2"
plain_sheet("Balance Issues", "Balance Issues")
plain_sheet("Change Log", "Change Log", drop=("Implementation Status",))

# ================================================================ Contents sheet (first)
ws = wb.create_sheet("Contents", 0)
ws.sheet_view.showGridLines = True
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 32
ws.column_dimensions["C"].width = 90
ws.merge_cells("B1:C1")
c = ws.cell(1, 2, "BOFA - GAME DOCUMENTATION")
c.font = Font(name=FONT, size=24, bold=True)
ws.row_dimensions[1].height = 45
c = ws.cell(2, 2, "Gen 1-6 Pokémon, moves and items only.  Orange cells = changed from vanilla.")
c.font = Font(name=FONT, color="666666")
DESC = {
    "Personal": "Base stats, types and abilities for every Gen 1-6 species (changed values in orange, with +/- vs vanilla)",
    "Ability Changes": "Only species whose abilities changed",
    "Evolutions": "Evolution method / requirement / result (changed evolutions in orange)",
    "Level-Up Learnsets": "Level-up learnset per species (Gen 1-6 moves); last column lists any Gen 7+ moves still in the game's learnset",
    "Moves": "Current data for every Gen 1-6 move (changed moves in orange)",
    "Move Changes": "One line per changed Gen 1-6 move: old -> new",
    "Encounters": "Wild encounters per map and method with % chance",
    "Gift Pokemon": "Gift, trade and one-time Pokemon by split and location: who gives it, what you give up, and its level, nature, ability, IVs, item and moves",
    "Split 1": "Restructured Split 1 in play order (in ROM, not playtested): 16 mandatory battles + 2 optional Viridian Forest trainers",
    "Split 1 Progression": "Step-by-step Split 1 route, gates and story flags",
    "Split 2 Progression": "Split 2 route and gates (woods open after Badge 1, Roxanne found on Route 116); teams not redesigned yet",
    "Split 3 Progression": "Split 3 route and trainer structure (Rusturf -> Wraithwood -> Hollowbrook -> Granite Cave -> Dewford); Brawly cap Lv34 tentative; teams TBD",
    "Split 1 Plan Check": "Trainer counts per area and every open decision (read this first)",
    "Items": "Every Gen 1-6 item: pocket, price, fling power, description and where to find it",
}
for i, n in enumerate(wb.sheetnames[1:], 4):
    c = ws.cell(i, 2, n)
    c.hyperlink = f"#'{n}'!A1"
    c.font = Font(name=FONT, bold=True, color="1155CC", underline="single")
    d = DESC.get(n, "Trainer teams in battle-card layout" if n.startswith(("Split", "Elite", "Post", "Unass", "Boss")) else "")
    ws.cell(i, 3, d).font = Font(name=FONT)

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT)
